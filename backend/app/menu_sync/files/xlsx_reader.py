"""XLSX → ``TabularData`` (story #65; ADR-8.2, ADR-8.7). Единственный модуль, импортирующий ``openpyxl``.

Зависимости: ``openpyxl>=3.1`` и ``defusedxml>=0.7`` (openpyxl сам использует defusedxml, если он
установлен, — защита от XML-бомб и внешних сущностей).

Безопасность до разбора (``check_xlsx_archive``, ``zipfile``, без openpyxl):

- сигнатура zip, читаемый каталог архива, наличие ``xl/workbook.xml``; ``vbaProject.bin``
  (макросы, ``.xlsm``, переименованный в ``.xlsx``) — ``unsupported_format``;
- сумма распакованных размеров ≤ ``MAX_UNCOMPRESSED_BYTES``, степень сжатия каждого члена ≤
  ``MAX_COMPRESSION_RATIO`` — иначе ``unsafe_archive`` (zip-бомба). Распаковка ``zipfile``
  не выходит за объявленный размер члена, поэтому проверка заголовков достаточна.

Разбор: ``load_workbook(BytesIO(data), read_only=True, data_only=True)`` — формулы не
вычисляются, берётся сохранённое значение. Вторая книга с ``data_only=False`` читается
параллельно построчно только затем, чтобы отличить формулу без сохранённого значения
(``formula_without_value``) от пустой ячейки. Лист — ``sheet`` из аргумента (точное имя, иначе
без учёта регистра), иначе лист ``menu`` (без учёта регистра), иначе первый. Ячейки с ошибками
Excel (``#N/A``) и длительностями — ``CellProblem`` (``invalid_cell_type``), даты — как есть
(их отвергает ``format``). Объявленные размеры листа (``<dimension>``) не используются: строки и
колонки берутся из самих ячеек. Номер строки — номер строки листа. Книги закрываются в ``finally``.

Лимиты — как в CSV: ``MAX_DATA_ROWS`` строк данных и ``MAX_COLUMNS`` колонок; ячейки без
значения правее ``MAX_COLUMNS`` (отформатированные, но пустые) отбрасываются, значение там —
``too_many_columns``.

Ограничение обхода (security review #65, DoS по CPU): openpyxl в ``read_only`` сам отдаёт пустую
строку для каждого пропущенного номера ``<row r=…>``, а строку с ячейкой в колонке XFD дополняет
до 16 384 элементов. Поэтому обход прерывается, если номер строки листа больше
``MAX_SHEET_ROWS`` (строки данных + ``MAX_BLANK_ROWS`` пустых, как ``MAX_BLANK_RECORDS`` в CSV), —
``too_many_rows``; или если сумма ширин просмотренных строк больше ``MAX_SCANNED_CELLS`` —
``too_many_columns``. Генераторы openpyxl ленивые: после ``break`` дальше они не идут.

Стоимость разбора openpyxl пропорциональна числу XML-элементов (десятки микросекунд на элемент,
а строки с повторным или убывающим ``r`` openpyxl разбирает, но не отдаёт). Поэтому до разбора
элементы считаются регулярным выражением (``_count_tags``, с ранним выходом):

- все части архива вне ``xl/worksheets/`` (общие строки, стили, свойства; их openpyxl читает при
  открытии обеих книг, имя части может быть любым) — не больше ``MAX_NON_SHEET_BYTES`` байт и
  ``MAX_NON_SHEET_TAGS`` элементов, вся книга — не больше ``MAX_BOOK_TAGS`` элементов, иначе
  ``unsafe_archive``;
- выбранный лист — не больше ``MAX_SHEET_ROWS`` элементов ``row``, ``MAX_SHEET_CELLS`` элементов
  ``c``, ``MAX_SHEET_OTHER_TAGS`` прочих элементов (объединения, условное форматирование…) и
  ``MAX_SHEET_TAGS`` всего.

Вторая книга (``data_only=False``) открывается, только если на листе есть элементы ``f``.
"""

from __future__ import annotations

import io
import re
import zipfile
from collections import Counter
from collections.abc import Iterator, Sequence
from datetime import timedelta
from pathlib import PurePosixPath
from typing import Any, Final

from openpyxl import load_workbook

from app.menu_sync.files.format import (
    MAX_COLUMNS,
    MAX_DATA_ROWS,
    REQUIRED_COLUMNS,
    Cell,
    CellProblem,
    SheetRow,
    TabularData,
)
from app.menu_sync.report import IssueCode, SyncIssue

DEFAULT_SHEET_NAME: Final[str] = "menu"
MAX_UNCOMPRESSED_BYTES: Final[int] = 50 * 1024 * 1024
MAX_COMPRESSION_RATIO: Final[int] = 100
# Части вне xl/worksheets/ (общие строки, стили…) читаются при открытии каждой из двух книг.
SHEETS_DIR: Final[str] = "xl/worksheets/"
MAX_NON_SHEET_BYTES: Final[int] = 20 * 1024 * 1024
MAX_NON_SHEET_TAGS: Final[int] = 200_000
# Пустые строки листа (в том числе досозданные openpyxl для пропусков номеров) — как в CSV.
MAX_BLANK_ROWS: Final[int] = MAX_DATA_ROWS
MAX_SHEET_ROWS: Final[int] = MAX_DATA_ROWS + MAX_BLANK_ROWS
# XML-элементы выбранного листа: ячейки (5000 строк по MAX_COLUMNS = 50 колонок), прочие и всего.
MAX_SHEET_CELLS: Final[int] = MAX_DATA_ROWS * MAX_COLUMNS
# Содержимое одной ячейки: значение и формула (v, f) или строка (is, t).
CELL_CONTENT_PER_CELL: Final[int] = 3
# Сумма ширин просмотренных строк (с дополнением openpyxl до последней ячейки строки):
# MAX_SHEET_ROWS строк по 50 колонок или около 30 строк с ячейкой в колонке XFD.
MAX_SCANNED_CELLS: Final[int] = 2 * MAX_SHEET_CELLS
MAX_SHEET_OTHER_TAGS: Final[int] = 20_000
MAX_SHEET_TAGS: Final[int] = (1 + CELL_CONTENT_PER_CELL) * MAX_SHEET_CELLS + MAX_SHEET_ROWS + MAX_SHEET_OTHER_TAGS
# Вся книга: лист меню и служебные части; другие большие листы в книгу с меню не кладутся.
MAX_BOOK_TAGS: Final[int] = MAX_SHEET_TAGS + MAX_NON_SHEET_TAGS
# Элементы содержимого ячейки (значение, формула, строка с форматированием); остальное — «прочие».
CELL_CONTENT_TAGS: Final[frozenset[bytes]] = frozenset(
    {b"v", b"f", b"is", b"t", b"r", b"rPr", b"rFont", b"b", b"i", b"u", b"strike", b"sz", b"color"}
    | {b"family", b"scheme", b"vertAlign", b"outline", b"shadow", b"condense", b"extend", b"charset", b"rPh"}
)
# Любая разметка, кроме закрывающего тега: «<» и сразу имя (пробел между ними XML не допускает).
# Префикс и имя — любые байты, кроме разделителей (NCName бывает не из ASCII: «<é:row>»); локальное
# имя — после префикса. «<?…», «<!--…», «<![CDATA[…» и «<» внутри CDATA или комментария считаются
# с пустым или случайным именем — как «прочие» элементы (запас в безопасную сторону).
_START_TAG: Final[re.Pattern[bytes]] = re.compile(rb"<(?!/)(?:[^\s/>!?:=<]+:)?([^\s/>!?:=<]*)")
# Коды ошибок Excel, которые можно показать в отчёте; иное значение ячейки t="e" не выводится.
EXCEL_ERROR_CODES: Final[frozenset[str]] = frozenset(
    {"#N/A", "#REF!", "#VALUE!", "#DIV/0!", "#NAME?", "#NUM!", "#NULL!", "#SPILL!", "#CALC!", "#GETTING_DATA"}
)

ZIP_SIGNATURE: Final[bytes] = b"PK\x03\x04"
WORKBOOK_PART: Final[str] = "xl/workbook.xml"
MACRO_PART_NAME: Final[str] = "vbaproject.bin"
ERROR_DATA_TYPE: Final[str] = "e"  # тип ячейки openpyxl «ошибка Excel»

NOT_XLSX_HINT: Final[str] = "файл не является книгой .xlsx или повреждён: откройте его в Excel и сохраните как .xlsx"
MACRO_HINT: Final[str] = "книга с макросами (.xlsm) не поддерживается: сохраните её как «Книга Excel (.xlsx)»"
UNREADABLE_HINT: Final[str] = "не удалось прочитать книгу .xlsx: откройте её в Excel и сохраните заново"
FORMULA_HINT: Final[str] = (
    "формула без сохранённого значения: откройте файл в Excel, дождитесь пересчёта и сохраните, "
    "или замените формулу значением"
)


def _archive_issue(code: IssueCode, message: str) -> list[SyncIssue]:
    return [SyncIssue(code, message)]


def _count_tags(xml: bytes, limit: int) -> Counter[bytes] | None:
    """Число открывающих тегов по локальному имени; ``None`` — тегов больше ``limit`` (счёт прерван)."""
    counts: Counter[bytes] = Counter()
    # Expat читает только ASCII-совместимые кодировки и UTF-16. В UTF-16 разметка — те же байты
    # ASCII через нулевой: без нулей теги видны так же, как в UTF-8 (в UTF-8 нулей в XML нет).
    for total, match in enumerate(_START_TAG.finditer(xml.replace(b"\x00", b"")), start=1):
        if total > limit:
            return None
        counts[match.group(1)] += 1
    return counts


def _is_sheet_part(name: str) -> bool:
    return name.startswith(SHEETS_DIR) and "/" not in name[len(SHEETS_DIR) :]


def check_xlsx_archive(data: bytes) -> list[SyncIssue]:
    """Проверки zip-контейнера до openpyxl (сигнатура, макросы, zip-бомба, объём служебных частей);
    ``[]`` — можно разбирать."""
    if not data.startswith(ZIP_SIGNATURE):
        return _archive_issue(IssueCode.UNSUPPORTED_FORMAT, NOT_XLSX_HINT)
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            members = archive.infolist()
            issues = _check_members(members)
            if not issues:
                issues = _check_tags(archive, members)
            return issues
    # RuntimeError — зашифрованный член, NotImplementedError — неизвестный метод сжатия.
    except (zipfile.BadZipFile, zipfile.LargeZipFile, ValueError, EOFError, OSError, RuntimeError, NotImplementedError):
        return _archive_issue(IssueCode.UNSUPPORTED_FORMAT, NOT_XLSX_HINT)


def _check_tags(archive: zipfile.ZipFile, members: list[zipfile.ZipInfo]) -> list[SyncIssue]:
    """Элементы разметки всех частей: вне листов — ``MAX_NON_SHEET_TAGS``, всего — ``MAX_BOOK_TAGS``
    (openpyxl при открытии читает начало каждого листа, до ``sheetData``)."""
    book_left, non_sheet_left = MAX_BOOK_TAGS, MAX_NON_SHEET_TAGS
    for member in members:
        is_sheet = _is_sheet_part(member.filename)
        counts = _count_tags(archive.read(member), book_left if is_sheet else min(book_left, non_sheet_left))
        if counts is None:
            limit = MAX_BOOK_TAGS if is_sheet or non_sheet_left > book_left else MAX_NON_SHEET_TAGS
            where = "книга" if limit == MAX_BOOK_TAGS else "служебные части книги (кроме листов)"
            message = f"{where} содержит больше {limit} элементов разметки: оставьте в книге только лист меню"
            return _archive_issue(IssueCode.UNSAFE_ARCHIVE, message)
        book_left -= counts.total()
        if not is_sheet:
            non_sheet_left -= counts.total()
    return []


def _check_members(members: list[zipfile.ZipInfo]) -> list[SyncIssue]:
    names = {member.filename for member in members}
    if any(PurePosixPath(name).name.lower() == MACRO_PART_NAME for name in names):
        return _archive_issue(IssueCode.UNSUPPORTED_FORMAT, MACRO_HINT)
    if WORKBOOK_PART not in names:
        return _archive_issue(IssueCode.UNSUPPORTED_FORMAT, NOT_XLSX_HINT)

    limit_mb = MAX_UNCOMPRESSED_BYTES // (1024 * 1024)
    if sum(member.file_size for member in members) > MAX_UNCOMPRESSED_BYTES:
        return _archive_issue(IssueCode.UNSAFE_ARCHIVE, f"распакованная книга больше {limit_mb} МБ")
    if sum(member.file_size for member in members if not _is_sheet_part(member.filename)) > MAX_NON_SHEET_BYTES:
        limit_mb = MAX_NON_SHEET_BYTES // (1024 * 1024)
        return _archive_issue(IssueCode.UNSAFE_ARCHIVE, f"служебные части книги (кроме листов) больше {limit_mb} МБ")
    for member in members:
        if member.file_size > MAX_COMPRESSION_RATIO * max(member.compress_size, 1):
            message = f"часть книги сжата больше чем в {MAX_COMPRESSION_RATIO} раз: файл похож на zip-бомбу"
            return _archive_issue(IssueCode.UNSAFE_ARCHIVE, message)
    return []


def _select_sheet(workbook: Any, sheet: str | None) -> tuple[Any | None, list[SyncIssue]]:
    """Лист по правилам модуля; ``None`` и ``sheet_not_found`` со списком листов, если его нет."""
    worksheets = list(workbook.worksheets)
    if sheet is not None:
        found = next((ws for ws in worksheets if ws.title == sheet), None)
        if found is None:
            folded = sheet.strip().casefold()
            found = next((ws for ws in worksheets if ws.title.casefold() == folded), None)
        if found is not None:
            return found, []
        available = ", ".join(ws.title for ws in worksheets) or "нет листов"
        message = f"в книге нет листа с таким именем; листы: {available}"
        return None, [SyncIssue(IssueCode.SHEET_NOT_FOUND, message, value=sheet)]
    if not worksheets:
        return None, [SyncIssue(IssueCode.SHEET_NOT_FOUND, "в книге нет листов с таблицей")]
    default = next((ws for ws in worksheets if ws.title.strip().casefold() == DEFAULT_SHEET_NAME), None)
    return default or worksheets[0], []


def _cell_value(cell: Any, formula: object) -> Cell:
    """Значение ячейки из книги ``data_only=True`` с учётом той же ячейки книги с формулами."""
    value = cell.value
    if value is None:
        # Пустая в data_only, но непустая без него — формула без сохранённого значения.
        return None if formula is None else CellProblem(IssueCode.FORMULA_WITHOUT_VALUE, FORMULA_HINT)
    if cell.data_type == ERROR_DATA_TYPE:
        shown = f"ошибка Excel {value}" if value in EXCEL_ERROR_CODES else "ошибка Excel"
        return CellProblem(IssueCode.INVALID_CELL_TYPE, f"{shown} вместо значения")
    if isinstance(value, timedelta):
        return CellProblem(IssueCode.INVALID_CELL_TYPE, "дата или время вместо значения: проверьте формат ячейки")
    return value


def _fit_columns(cells: Sequence[Any], formulas: Sequence[object]) -> bool:
    """Помещается ли строка в ``MAX_COLUMNS``: правее лимита допустимы только ячейки без значения.

    ``count`` — без цикла Python: отформатированная пустая ячейка в колонке XFD даёт строку
    из 16 384 элементов.
    """
    tail = formulas[MAX_COLUMNS:]
    return len(cells) <= MAX_COLUMNS or tail.count(None) == len(tail)


def _is_blank(values: Sequence[Cell]) -> bool:
    return all(value is None or (isinstance(value, str) and not value.strip()) for value in values)


def _header_text(value: Cell, position: int, row: int) -> str | SyncIssue:
    if isinstance(value, CellProblem):
        return SyncIssue(value.code, value.message, row=row, column=f"колонка {position + 1}")
    if value is None:
        return ""
    return value if isinstance(value, str) else str(value)


def _check_sheet_xml(xml: bytes) -> tuple[list[SyncIssue], bool]:
    """Число XML-элементов выбранного листа до разбора openpyxl; второе значение — есть ли формулы."""
    counts = _count_tags(xml, MAX_SHEET_TAGS)
    if counts is None:
        message = f"лист содержит больше {MAX_SHEET_TAGS} элементов разметки: оставьте на листе только таблицу меню"
        return [SyncIssue(IssueCode.TOO_MANY_ROWS, message)], False
    if counts[b"row"] > MAX_SHEET_ROWS:
        message = f"на листе больше {MAX_SHEET_ROWS} строк вместе с пустыми: удалите строки ниже таблицы"
        return [SyncIssue(IssueCode.TOO_MANY_ROWS, message)], False
    if counts[b"c"] > MAX_SHEET_CELLS:
        message = f"на листе больше {MAX_SHEET_CELLS} ячеек вместе с пустыми отформатированными"
        return [SyncIssue(IssueCode.TOO_MANY_COLUMNS, message)], False
    # openpyxl разбирает как ячейку любой дочерний элемент row: содержимое ячеек сверх
    # CELL_CONTENT_PER_CELL на элемент c считается «прочими» элементами.
    content = sum(counts[name] for name in CELL_CONTENT_TAGS)
    excess = max(0, content - CELL_CONTENT_PER_CELL * counts[b"c"])
    other = counts.total() - counts[b"row"] - counts[b"c"] - content + excess
    if other > MAX_SHEET_OTHER_TAGS:
        message = (
            f"на листе больше {MAX_SHEET_OTHER_TAGS} служебных элементов "
            "(объединения ячеек, условное форматирование и т. п.): скопируйте таблицу на новый лист"
        )
        return [SyncIssue(IssueCode.UNSAFE_ARCHIVE, message)], False
    return [], counts[b"f"] > 0


def _sheet_xml(data: bytes, worksheet: Any) -> bytes | None:
    """XML выбранного листа из архива (путь — из openpyxl); ``None`` — путь неизвестен."""
    path = getattr(worksheet, "_worksheet_path", None)
    if not isinstance(path, str):
        return None
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        return archive.read(path)


def _rows(worksheet: Any, formula_sheet: Any | None) -> Iterator[tuple[int, Sequence[Any], Sequence[object] | None]]:
    """Строки листа (номер с 1, ячейки ``data_only``, значения с формулами или ``None`` — формул
    на листе нет) без объявленных размеров."""
    worksheet.reset_dimensions()
    if formula_sheet is None:
        for number, cells in enumerate(worksheet.iter_rows(), start=1):
            yield number, cells, None
        return
    formula_sheet.reset_dimensions()
    pairs = zip(worksheet.iter_rows(), formula_sheet.iter_rows(values_only=True), strict=True)
    for number, (cells, formulas) in enumerate(pairs, start=1):
        yield number, cells, formulas


def _read_sheet(worksheet: Any, formula_sheet: Any | None) -> tuple[TabularData | None, list[SyncIssue]]:
    header: tuple[str, ...] | None = None
    header_row = 0
    rows: list[SheetRow] = []
    scanned_cells = 0
    for number, cells, formula_values in _rows(worksheet, formula_sheet):
        if number > MAX_SHEET_ROWS:
            message = (
                f"на листе больше {MAX_SHEET_ROWS} строк вместе с пустыми: удалите пустые строки и ячейки ниже таблицы"
            )
            return None, [SyncIssue(IssueCode.TOO_MANY_ROWS, message, row=number)]
        scanned_cells += len(cells)
        if scanned_cells > MAX_SCANNED_CELLS:
            message = "лист слишком широкий: удалите отформатированные пустые ячейки правее таблицы"
            return None, [SyncIssue(IssueCode.TOO_MANY_COLUMNS, message, row=number)]
        # Без формул на листе значение «с формулами» совпадает со значением data_only.
        formulas = formula_values if formula_values is not None else tuple(cell.value for cell in cells)
        if not _fit_columns(cells, formulas):
            return None, [SyncIssue(IssueCode.TOO_MANY_COLUMNS, f"больше {MAX_COLUMNS} колонок", row=number)]
        values = tuple(
            _cell_value(cell, formula)
            for cell, formula in zip(cells[:MAX_COLUMNS], formulas[:MAX_COLUMNS], strict=True)
        )
        if _is_blank(values):
            continue
        if header is None:
            texts = [_header_text(value, position, number) for position, value in enumerate(values)]
            problems = [text for text in texts if isinstance(text, SyncIssue)]
            if problems:
                return None, problems
            header, header_row = tuple(str(text) for text in texts), number
            continue
        if len(rows) == MAX_DATA_ROWS:
            return None, [SyncIssue(IssueCode.TOO_MANY_ROWS, f"больше {MAX_DATA_ROWS} строк данных", row=number)]
        rows.append(SheetRow(number=number, values=values))

    if header is None:
        missing = ", ".join(REQUIRED_COLUMNS)
        message = f"лист пуст: нет строки заголовка с колонками {missing}"
        return None, [SyncIssue(IssueCode.EMPTY_SNAPSHOT, message, value=worksheet.title)]
    return TabularData(header=header, header_row=header_row, rows=rows), []


def read_xlsx_table(data: bytes, *, sheet: str | None = None) -> tuple[TabularData | None, list[SyncIssue]]:
    """Прочитать лист XLSX из байтов; ``None`` и проблемы уровня файла/листа/заголовка при неудаче."""
    issues = check_xlsx_archive(data)
    if issues:
        return None, issues
    workbook: Any = None
    formula_book: Any = None
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        worksheet, issues = _select_sheet(workbook, sheet)
        if worksheet is None:
            return None, issues
        xml = _sheet_xml(data, worksheet)
        if xml is None:
            return None, [SyncIssue(IssueCode.UNSUPPORTED_FORMAT, UNREADABLE_HINT)]
        issues, has_formulas = _check_sheet_xml(xml)
        if issues:
            return None, issues
        if not has_formulas:
            return _read_sheet(worksheet, None)
        formula_book = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
        return _read_sheet(worksheet, formula_book[worksheet.title])
    except Exception:  # недоверенный файл: openpyxl и XML-парсер бросают исключения разных типов
        return None, [SyncIssue(IssueCode.UNSUPPORTED_FORMAT, UNREADABLE_HINT)]
    finally:
        for book in (workbook, formula_book):
            if book is not None:
                book.close()
