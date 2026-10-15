"""Файл меню XLSX → снимок (story #65; ADR-8.2, ADR-8.7): ``app.menu_sync.files.xlsx_reader``, без БД.

AC #65: книга, где первый лист — другой, а на втором листе ``Menu`` — данные синтетического
CSV-примера с ценами-числами; разбор без указания листа выбирает ``Menu``, снимок совпадает со
снимком ``menu_valid.csv``, цены — ``Decimal`` с двумя знаками без двоичной погрешности.

Книги собираются в самих тестах (openpyxl); данные синтетические.
"""

from __future__ import annotations

import csv
import io
import math
import time
import zipfile
from collections.abc import Callable
from datetime import date, datetime, timedelta
from decimal import Decimal
from pathlib import Path
from typing import Any, ClassVar

import openpyxl
import openpyxl.xml
import pytest
from defusedxml.common import DefusedXmlException
from openpyxl import Workbook

from app.menu_sync.files import FileFormat, ParsedMenuFile, parse_menu_bytes, read_menu_file, xlsx_reader
from app.menu_sync.files.format import MAX_COLUMNS
from app.menu_sync.files.xlsx_reader import (
    MACRO_HINT,
    MAX_COMPRESSION_RATIO,
    NOT_XLSX_HINT,
    UNREADABLE_HINT,
    check_xlsx_archive,
    read_xlsx_table,
)
from app.menu_sync.report import MAX_REPORTED_ISSUES, IssueCode, SyncIssue

FIXTURES = Path(__file__).resolve().parent / "fixtures" / "menu"
VALID_CSV = FIXTURES / "menu_valid.csv"

HEADER = ["category_external_id", "category_name", "external_id", "name", "price_default"]

Rows = list[list[Any]]


def book(*sheets: tuple[str, Rows], cells: dict[str, dict[str, Any]] | None = None) -> bytes:
    """Книга из листов (имя, строки); ``cells`` — отдельные ячейки по листам (``{"Menu": {"A5": 1}}``)."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, rows in sheets:
        worksheet = workbook.create_sheet(title)
        for row in rows:
            worksheet.append(row)
        for coordinate, value in (cells or {}).get(title, {}).items():
            worksheet[coordinate] = value
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def menu_book(*rows: list[Any], header: list[str] | None = None) -> bytes:
    return book(("menu", [header or HEADER, *rows]))


def rewrite(
    data: bytes, transform: Callable[[str, bytes], bytes | None], extra: dict[str, bytes] | None = None
) -> bytes:
    """Пересобрать zip: ``transform(name, content)`` → новое содержимое или ``None`` (удалить член)."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            content = transform(info.filename, source.read(info))
            if content is not None:
                target.writestr(info.filename, content)
        for name, content in (extra or {}).items():
            target.writestr(name, content)
    return buffer.getvalue()


def with_cached_value(data: bytes, formula: str, cached: str) -> bytes:
    """Дописать сохранённое значение формуле (openpyxl пишет формулы без него, как библиотека без пересчёта)."""

    def transform(name: str, content: bytes) -> bytes:
        if not name.startswith("xl/worksheets/"):
            return content
        return content.replace(f"<f>{formula}</f><v />".encode(), f"<f>{formula}</f><v>{cached}</v>".encode())

    return rewrite(data, transform)


def parse(data: bytes, sheet: str | None = None) -> ParsedMenuFile:
    return parse_menu_bytes(data, "menu.xlsx", sheet=sheet)


def codes(result: ParsedMenuFile) -> list[IssueCode]:
    return [issue.code for issue in result.issues]


def only_issue(result: ParsedMenuFile) -> SyncIssue:
    assert result.snapshot is None
    assert len(result.issues) == 1, [issue.to_dict() for issue in result.issues]
    return result.issues[0]


def price_of(result: ParsedMenuFile, external_id: str, tier: str = "default") -> Decimal:
    assert result.snapshot is not None, [issue.to_dict() for issue in result.issues]
    return next(item for item in result.snapshot.items if item.external_id == external_id).prices[tier]


# ---------------------------------------------------------------- AC #65


def _typed(column: str, text: str) -> Any:
    """Значение CSV-примера так, как его хранит Excel: числа — числами, id с ведущим нулём — текстом."""
    if text == "":
        return None
    if column.startswith("price_"):
        return float(text)
    if column in {"category_external_id", "category_sort_key", "sort_key"}:
        return int(text)
    if column == "external_id":
        return text if text.startswith("0") else int(text)
    if column == "is_available":
        return {"1": 1, "0": 0, "true": True}.get(text, text)
    return text


def csv_example_rows() -> Rows:
    with VALID_CSV.open(encoding="utf-8", newline="") as stream:
        header, *records = list(csv.reader(stream))
    return [
        header,
        *([_typed(column, text) for column, text in zip(header, record, strict=True)] for record in records),
    ]


def test_ac65_menu_sheet_matches_csv_example() -> None:
    rows = csv_example_rows()
    assert any(isinstance(value, float) for value in rows[1])  # цены действительно числа
    data = book(("Prices", [["note"], ["другой лист"]]), ("Menu", rows))

    result = parse(data)
    expected = read_menu_file(VALID_CSV)

    assert result.issues == []
    assert result.file_format is FileFormat.XLSX
    assert result.snapshot is not None
    assert result.snapshot == expected.snapshot
    assert result.item_rows == expected.item_rows
    assert result.rows_total == expected.rows_total == 9
    assert result.has_availability is True
    for item in result.snapshot.items:
        for price in item.prices.values():
            assert isinstance(price, Decimal)
            assert price.as_tuple().exponent == -2
    assert price_of(result, "007") == Decimal("4.50")
    assert str(price_of(result, "103")) == "7.80"
    assert str(price_of(result, "008")) == "6.00"


def test_defusedxml_is_used_by_openpyxl() -> None:
    assert openpyxl.xml.DEFUSEDXML is True


# ---------------------------------------------------------------- выбор листа


def two_sheets(first: str = "First", second: str = "Second") -> bytes:
    return book(
        (first, [HEADER, ["10", "A", "1", "из первого", 1]]),
        (second, [HEADER, ["10", "A", "2", "из второго", 2]]),
    )


def item_ids(result: ParsedMenuFile) -> list[str]:
    assert result.snapshot is not None, [issue.to_dict() for issue in result.issues]
    return [item.external_id for item in result.snapshot.items]


def test_without_menu_sheet_first_is_used() -> None:
    assert item_ids(parse(two_sheets())) == ["1"]


@pytest.mark.parametrize("title", ["menu", "MENU", "Menu ", "mEnU"])
def test_menu_sheet_name_is_case_insensitive(title: str) -> None:
    assert item_ids(parse(two_sheets(second=title))) == ["2"]


@pytest.mark.parametrize("sheet", ["Second", "second", " SECOND "])
def test_sheet_argument_exact_then_case_insensitive(sheet: str) -> None:
    assert item_ids(parse(two_sheets(second="Second"), sheet=sheet)) == ["2"]


def test_sheet_argument_overrides_menu_sheet() -> None:
    data = book(
        ("menu", [HEADER, ["10", "A", "1", "menu", 1]]),
        ("Data", [HEADER, ["10", "A", "2", "Data", 2]]),
    )
    assert item_ids(parse(data, sheet="Data")) == ["2"]
    assert item_ids(parse(data)) == ["1"]


def test_sheet_argument_exact_name_wins_over_case_insensitive() -> None:
    """Excel не допускает имён листов, отличающихся только регистром; правило проверяется на ``_select_sheet``."""

    class Sheet:
        def __init__(self, title: str) -> None:
            self.title = title

    class Book:
        worksheets: ClassVar[list[Sheet]] = [Sheet("Data"), Sheet("data")]

    assert xlsx_reader._select_sheet(Book(), "data")[0] is Book.worksheets[1]
    assert xlsx_reader._select_sheet(Book(), "DATA")[0] is Book.worksheets[0]


def test_sheet_not_found_lists_sheets() -> None:
    issue = only_issue(parse(two_sheets(), sheet="Меню"))
    assert issue.code is IssueCode.SHEET_NOT_FOUND
    assert issue.value == "Меню"
    assert "First, Second" in issue.message


def test_workbook_without_worksheets(monkeypatch: pytest.MonkeyPatch) -> None:
    """Книга только из диаграмм: openpyxl отдаёт пустой ``worksheets``."""

    class Empty:
        worksheets: ClassVar[list[Any]] = []

    assert xlsx_reader._select_sheet(Empty(), None)[1][0].code is IssueCode.SHEET_NOT_FOUND
    issue = xlsx_reader._select_sheet(Empty(), "x")[1][0]
    assert issue.code is IssueCode.SHEET_NOT_FOUND
    assert "нет листов" in issue.message


# ---------------------------------------------------------------- типы ячеек


def test_numeric_external_id_without_fraction() -> None:
    result = parse(menu_book([10.0, "A", 7.0, "Dolma", 8]))
    assert item_ids(result) == ["7"]
    assert result.snapshot is not None
    assert result.snapshot.categories[0].external_id == "10"


def test_fractional_external_id_is_invalid_id() -> None:
    issue = only_issue(parse(menu_book(["10", "A", 7.5, "Dolma", 8])))
    assert (issue.code, issue.row, issue.column) == (IssueCode.INVALID_ID, 2, "external_id")


def test_text_external_id_keeps_leading_zeros() -> None:
    assert item_ids(parse(menu_book(["10", "A", "007", "Dolma", 8]))) == ["007"]


@pytest.mark.parametrize(
    ("cell", "expected"),
    [(8, "8.00"), (8.5, "8.50"), (8.1, "8.10"), (0.07, "0.07"), (99999999.99, "99999999.99"), ("4.50", "4.50")],
)
def test_numeric_prices_are_exact(cell: Any, expected: str) -> None:
    assert str(price_of(parse(menu_book(["10", "A", "1", "Dolma", cell])), "1")) == expected


def test_float_noise_from_formula_is_rounded() -> None:
    data = with_cached_value(menu_book(["10", "A", "1", "Dolma", "=0.1+0.2"]), "0.1+0.2", "0.30000000000000004")
    assert str(price_of(parse(data), "1")) == "0.30"


@pytest.mark.parametrize(
    ("cell", "code"),
    [
        (8.123, IssueCode.TOO_MANY_DECIMALS),
        (-1.5, IssueCode.NEGATIVE_PRICE),
        (True, IssueCode.INVALID_MONEY),
        ("8,50", IssueCode.INVALID_MONEY),  # десятичная запятая только в CSV с «;»
        (datetime(2024, 1, 2, 3, 4), IssueCode.INVALID_CELL_TYPE),  # noqa: DTZ001 — в Excel нет часовых поясов
        (date(2024, 1, 2), IssueCode.INVALID_CELL_TYPE),
        (timedelta(hours=1), IssueCode.INVALID_CELL_TYPE),
        ("#N/A", IssueCode.INVALID_CELL_TYPE),
        ("#DIV/0!", IssueCode.INVALID_CELL_TYPE),
    ],
)
def test_price_cell_errors(cell: Any, code: IssueCode) -> None:
    issue = only_issue(parse(menu_book(["10", "A", "1", "Dolma", cell])))
    assert (issue.code, issue.row, issue.column, issue.external_id) == (code, 2, "price_default", "1")


def test_excel_error_message_names_error() -> None:
    issue = only_issue(parse(menu_book(["10", "A", "1", "#N/A", 8])))
    assert (issue.code, issue.column) == (IssueCode.INVALID_CELL_TYPE, "name")
    assert "#N/A" in issue.message


def with_error_value(data: bytes, old: str, new: str) -> bytes:
    def transform(name: str, content: bytes) -> bytes:
        if not name.startswith("xl/worksheets/"):
            return content
        return content.replace(f"<v>{old}</v>".encode(), f"<v>{new}</v>".encode())

    return rewrite(data, transform)


def test_unknown_excel_error_value_is_not_shown() -> None:
    """Security review #65 [LOW]: значение ячейки t="e" из файла показывается только из белого списка."""
    data = with_error_value(menu_book(["10", "A", "1", "#N/A", 8]), "#N/A", "#" + "X" * 500)
    issue = only_issue(parse(data))
    assert (issue.code, issue.column) == (IssueCode.INVALID_CELL_TYPE, "name")
    assert issue.message == "ошибка Excel вместо значения"


@pytest.mark.parametrize("code", sorted(xlsx_reader.EXCEL_ERROR_CODES))
def test_known_excel_errors_are_shown(code: str) -> None:
    data = with_error_value(menu_book(["10", "A", "1", "#N/A", 8]), "#N/A", code)
    assert only_issue(parse(data)).message == f"ошибка Excel {code} вместо значения"


def test_date_in_external_id_is_invalid_id() -> None:
    moment = datetime(2024, 1, 2)  # noqa: DTZ001 — в Excel нет часовых поясов
    issue = only_issue(parse(menu_book(["10", "A", moment, "Dolma", 8])))
    assert (issue.code, issue.column) == (IssueCode.INVALID_ID, "external_id")


def test_excel_error_in_external_id_keeps_its_code() -> None:
    issue = only_issue(parse(menu_book(["10", "A", "#REF!", "Dolma", 8])))
    assert (issue.code, issue.column) == (IssueCode.INVALID_CELL_TYPE, "external_id")


def test_bool_and_numbers_in_is_available() -> None:
    header = [*HEADER, "is_available"]
    data = menu_book(
        ["10", "A", "1", "a", 1, True],
        ["10", "A", "2", "b", 1, False],
        ["10", "A", "3", "c", 1, 1],
        ["10", "A", "4", "d", 1, 0],
        header=header,
    )
    result = parse(data)
    assert result.snapshot is not None
    assert [item.is_available for item in result.snapshot.items] == [True, False, True, False]


def test_formula_with_cached_value_uses_value() -> None:
    data = menu_book(["10", "A", "1", "Dolma", "=4+4.5"])
    assert str(price_of(parse(with_cached_value(data, "4+4.5", "8.5")), "1")) == "8.50"


def test_text_formula_with_cached_value() -> None:
    data = with_cached_value(menu_book(["10", "A", "1", '=UPPER("dolma")', 8]), 'UPPER("dolma")', "DOLMA")
    # Excel помечает текстовый результат формулы типом t="str".
    data = rewrite(data, lambda _, content: content.replace(b'<c r="D2">', b'<c r="D2" t="str">'))
    result = parse(data)
    assert result.snapshot is not None
    assert result.snapshot.items[0].name == "DOLMA"


def test_formula_without_value_is_reported_with_row_and_column() -> None:
    result = parse(menu_book(["10", "A", "1", "Dolma", 8], ["10", "A", "2", "Kabab", "=B3*2"]))
    issue = only_issue(result)
    assert (issue.code, issue.row, issue.column, issue.external_id) == (
        IssueCode.FORMULA_WITHOUT_VALUE,
        3,
        "price_default",
        "2",
    )
    assert "Excel" in issue.message


def test_formula_without_value_in_comment_column_is_ignored() -> None:
    data = menu_book(["10", "A", "1", "Dolma", 8, "=E2*2"], header=[*HEADER, "# заметка"])
    assert str(price_of(parse(data), "1")) == "8.00"


def test_value_in_column_without_header_is_error() -> None:
    issue = only_issue(parse(menu_book(["10", "A", "1", "Dolma", 8, None, "лишнее"])))
    assert (issue.code, issue.column) == (IssueCode.UNKNOWN_COLUMN, "колонка 7")


# ---------------------------------------------------------------- заголовок и строки


def test_header_after_blank_rows_and_gaps_keep_sheet_row_numbers() -> None:
    data = book(
        ("menu", []),
        cells={
            "menu": {
                **{f"{letter}3": name for letter, name in zip("ABCDE", HEADER, strict=True)},
                "A4": "10",
                "B4": "A",
                "C4": "1",
                "D4": "Dolma",
                "E4": 8,
                "A5": " ",  # строка из пробелов пропускается
                "A9": "10",
                "B9": "A",
                "C9": "1",
                "D9": "Dublikat",
                "E9": 9,
            }
        },
    )
    issue = only_issue(parse(data))
    assert (issue.code, issue.row) == (IssueCode.DUPLICATE_EXTERNAL_ID, 9)
    assert "строка 4" in issue.message


def test_header_cell_with_excel_error() -> None:
    issue = only_issue(parse(menu_book(["10", "A", "1", "Dolma", 8], header=["category_external_id", "#N/A"])))
    assert (issue.code, issue.row, issue.column) == (IssueCode.INVALID_CELL_TYPE, 1, "колонка 2")


def test_header_formula_without_value() -> None:
    issue = only_issue(parse(menu_book(header=[*HEADER, '="x"'])))
    assert (issue.code, issue.row, issue.column) == (IssueCode.FORMULA_WITHOUT_VALUE, 1, "колонка 6")


def test_numeric_header_cell_is_unknown_column() -> None:
    result = parse(menu_book(["10", "A", "1", "Dolma", 8, None], header=[*HEADER, 7]))
    assert codes(result) == [IssueCode.UNKNOWN_COLUMN]
    assert result.issues[0].column == "7"


def test_header_with_empty_cells_between_columns() -> None:
    header = [HEADER[0], None, *HEADER[1:]]
    result = parse(menu_book(["10", None, "A", "1", "Dolma", 8], header=header))
    assert item_ids(result) == ["1"]


def test_empty_sheet() -> None:
    issue = only_issue(parse(book(("menu", []))))
    assert issue.code is IssueCode.EMPTY_SNAPSHOT
    assert issue.value == "menu"


def test_header_only() -> None:
    assert codes(parse(menu_book())) == [IssueCode.EMPTY_SNAPSHOT]


def test_too_many_rows(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(xlsx_reader, "MAX_DATA_ROWS", 2)
    rows = [["10", "A", str(n), "Dolma", 1] for n in range(3)]
    issue = only_issue(parse(menu_book(*rows)))
    assert (issue.code, issue.row) == (IssueCode.TOO_MANY_ROWS, 4)
    assert item_ids(parse(menu_book(*rows[:2]))) == ["0", "1"]


def add_rows(data: bytes, rows_xml: bytes) -> bytes:
    """Дописать в конец ``sheetData`` первого листа готовые элементы ``<row>``."""

    def transform(name: str, content: bytes) -> bytes:
        if name != "xl/worksheets/sheet1.xml":
            return content
        return content.replace(b"</sheetData>", rows_xml + b"</sheetData>")

    return rewrite(data, transform)


def timed_parse(data: bytes) -> tuple[ParsedMenuFile, float]:
    started = time.perf_counter()
    result = parse(data)
    return result, time.perf_counter() - started


@pytest.mark.parametrize("number", [10**7, 10**12])
def test_huge_row_number_stops_quickly(number: int) -> None:
    """Security review #65 [HIGH]: openpyxl досоздаёт пустые строки до номера ``r``."""
    row = f'<row r="{number}"><c r="A{number}" t="inlineStr"><is><t>x</t></is></c></row>'.encode()
    data = add_rows(menu_book(["10", "A", "1", "Dolma", 8]), row)
    assert len(data) < 10 * 1024
    result, elapsed = timed_parse(data)
    issue = only_issue(result)
    assert (issue.code, issue.row) == (IssueCode.TOO_MANY_ROWS, xlsx_reader.MAX_SHEET_ROWS + 1)
    assert elapsed < 5  # с запасом: на деле доли секунды, до исправления — десятки секунд и больше


def test_many_blank_formatted_rows_stop_before_parsing() -> None:
    """Строки сверх ``MAX_SHEET_ROWS`` отсекаются подсчётом элементов ``row`` до openpyxl."""
    last = xlsx_reader.MAX_SHEET_ROWS + 500
    rows = b"".join(f'<row r="{n}"><c r="B{n}" s="0" /></row>'.encode() for n in range(3, last))
    result, elapsed = timed_parse(add_rows(menu_book(["10", "A", "1", "Dolma", 8]), rows))
    issue = only_issue(result)
    assert issue.code is IssueCode.TOO_MANY_ROWS
    assert str(xlsx_reader.MAX_SHEET_ROWS) in issue.message
    assert elapsed < 5


def test_rows_with_repeated_number_are_counted_before_parsing() -> None:
    """openpyxl разбирает, но не отдаёт строки с повторным ``r``: обход их не видит, подсчёт — видит."""
    count = xlsx_reader.MAX_SHEET_ROWS + 1
    rows = b"".join(f'<row r="2" spans="{n}:{n}"><c r="XFD2" s="0" /></row>'.encode() for n in range(count))
    result, elapsed = timed_parse(add_rows(menu_book(["10", "A", "1", "Dolma", 8]), rows))
    assert codes(result) == [IssueCode.TOO_MANY_ROWS]
    assert elapsed < 5


def test_many_rows_with_formatted_cell_in_xfd_stop_at_cell_budget() -> None:
    rows = b"".join(f'<row r="{n}"><c r="XFD{n}" s="0" /></row>'.encode() for n in range(3, 5000))
    result, elapsed = timed_parse(add_rows(menu_book(["10", "A", "1", "Dolma", 8]), rows))
    issue = only_issue(result)
    assert issue.code is IssueCode.TOO_MANY_COLUMNS
    assert issue.row is not None and issue.row < 30  # примерно MAX_SCANNED_CELLS / 16 384
    assert "широкий" in issue.message
    assert elapsed < 5


def test_sheet_row_limit_is_inclusive(monkeypatch: pytest.MonkeyPatch) -> None:
    # три элемента row: заголовок, строки 2 и 5 листа (3 и 4 — пропуск номеров)
    row = ["10", "A", "2", "Kabab", 9]
    cells = {f"{letter}5": value for letter, value in zip("ABCDE", row, strict=True)}
    data = book(("menu", [HEADER, ["10", "A", "1", "Dolma", 8]]), cells={"menu": cells})
    monkeypatch.setattr(xlsx_reader, "MAX_SHEET_ROWS", 5)
    assert item_ids(parse(data)) == ["1", "2"]
    monkeypatch.setattr(xlsx_reader, "MAX_SHEET_ROWS", 4)
    issue = only_issue(parse(data))
    assert (issue.code, issue.row) == (IssueCode.TOO_MANY_ROWS, 5)
    monkeypatch.setattr(xlsx_reader, "MAX_SHEET_ROWS", 2)  # элементов row больше лимита
    issue = only_issue(parse(data))
    assert (issue.code, issue.row) == (IssueCode.TOO_MANY_ROWS, None)


@pytest.mark.parametrize(
    ("limit", "value", "code"),
    [
        ("MAX_SHEET_CELLS", 9, IssueCode.TOO_MANY_COLUMNS),  # 10 элементов c
        ("MAX_SHEET_TAGS", 20, IssueCode.TOO_MANY_ROWS),
        ("MAX_SHEET_OTHER_TAGS", 3, IssueCode.UNSAFE_ARCHIVE),  # worksheet, sheetPr, dimension, …
    ],
)
def test_sheet_xml_limits(monkeypatch: pytest.MonkeyPatch, limit: str, value: int, code: IssueCode) -> None:
    data = menu_book(["10", "A", "1", "Dolma", 8])
    monkeypatch.setattr(xlsx_reader, limit, value)
    assert codes(parse(data)) == [code]


def test_many_merged_cells_rejected_before_parsing() -> None:
    merges = "".join(f'<mergeCell ref="A{n}:B{n}"/>' for n in range(10, 10 + xlsx_reader.MAX_SHEET_OTHER_TAGS))
    xml = f"<mergeCells>{merges}</mergeCells>".encode()

    def transform(name: str, content: bytes) -> bytes:
        if name != "xl/worksheets/sheet1.xml":
            return content
        return content.replace(b"</sheetData>", b"</sheetData>" + xml)

    result, elapsed = timed_parse(rewrite(menu_book(["10", "A", "1", "Dolma", 8]), transform))
    issue = only_issue(result)
    assert issue.code is IssueCode.UNSAFE_ARCHIVE
    assert "служебных элементов" in issue.message
    assert elapsed < 5


def test_count_tags_ignores_closing_and_prefixes() -> None:
    counts = xlsx_reader._count_tags(b'<?xml version="1.0"?><x:row r="1"><c/><!-- <c --></x:row>', 10)
    assert counts == {b"row": 1, b"c": 2}  # тег в комментарии тоже считается: запас в безопасную сторону
    assert xlsx_reader._count_tags(b"<a><b/><c/></a>", 2) is None
    assert xlsx_reader._count_tags(b"<a><b/></a>", 2) == {b"a": 1, b"b": 1}


def test_sheet_path_unknown_is_unreadable(monkeypatch: pytest.MonkeyPatch) -> None:
    assert xlsx_reader._sheet_xml(b"", object()) is None
    monkeypatch.setattr(xlsx_reader, "_sheet_xml", lambda *_: None)
    issue = only_issue(parse(menu_book(["10", "A", "1", "Dolma", 8])))
    assert (issue.code, issue.message) == (IssueCode.UNSUPPORTED_FORMAT, UNREADABLE_HINT)


def test_scanned_cells_limit_is_inclusive(monkeypatch: pytest.MonkeyPatch) -> None:
    data = menu_book(["10", "A", "1", "Dolma", 8])  # две строки по 5 ячеек
    monkeypatch.setattr(xlsx_reader, "MAX_SCANNED_CELLS", 10)
    assert item_ids(parse(data)) == ["1"]
    monkeypatch.setattr(xlsx_reader, "MAX_SCANNED_CELLS", 9)
    issue = only_issue(parse(data))
    assert (issue.code, issue.row) == (IssueCode.TOO_MANY_COLUMNS, 2)


def test_sheet_limits_follow_csv() -> None:
    assert xlsx_reader.MAX_BLANK_ROWS == xlsx_reader.MAX_DATA_ROWS
    assert xlsx_reader.MAX_SHEET_ROWS == 2 * xlsx_reader.MAX_DATA_ROWS


def test_value_beyond_max_columns() -> None:
    row = ["10", "A", "1", "Dolma", 8, *([None] * (MAX_COLUMNS - 5)), "x"]
    assert len(row) == MAX_COLUMNS + 1
    issue = only_issue(parse(menu_book(row)))
    assert (issue.code, issue.row) == (IssueCode.TOO_MANY_COLUMNS, 2)


def test_formatted_empty_cell_far_right_is_ignored() -> None:
    """Отформатированная пустая ячейка в последней колонке листа не делает строку «широкой»."""

    def style_far_cell(name: str, content: bytes) -> bytes:
        if name != "xl/worksheets/sheet1.xml":
            return content
        rows = content.split(b"</row>")
        rows[1] += b'<c r="XFD2" s="0" />'
        return b"</row>".join(rows)

    data = rewrite(menu_book(["10", "A", "1", "Dolma", 8]), style_far_cell)
    assert b"XFD2" in zipfile.ZipFile(io.BytesIO(data)).read("xl/worksheets/sheet1.xml")
    assert item_ids(parse(data)) == ["1"]


def test_many_row_errors_are_limited() -> None:
    rows = [["10", "A", str(n), "Dolma", "#N/A"] for n in range(MAX_REPORTED_ISSUES + 20)]
    result = parse(menu_book(*rows))
    assert len(result.issues) == MAX_REPORTED_ISSUES
    assert result.issues_truncated == 20
    assert {issue.code for issue in result.issues} == {IssueCode.INVALID_CELL_TYPE}


def test_read_xlsx_table_returns_typed_values() -> None:
    table, issues = read_xlsx_table(menu_book(["10", "A", 7, "Dolma", 8.5, None]))
    assert issues == []
    assert table is not None
    assert (table.header, table.header_row, table.decimal_comma) == (tuple(HEADER), 1, False)
    assert [(row.number, row.values) for row in table.rows] == [(2, ("10", "A", 7, "Dolma", 8.5))]


def test_read_menu_file_xlsx_from_disk_with_sheet(tmp_path: Path) -> None:
    path = tmp_path / "Menu.XLSX"
    path.write_bytes(two_sheets())
    result = read_menu_file(path, sheet="Second")
    assert result.file_format is FileFormat.XLSX
    assert item_ids(result) == ["2"]


# ---------------------------------------------------------------- архив: безопасность до openpyxl


def zip_of(members: dict[str, bytes], compression: int = zipfile.ZIP_DEFLATED) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression) as archive:
        for name, content in members.items():
            archive.writestr(name, content)
    return buffer.getvalue()


def test_valid_book_passes_archive_check() -> None:
    assert check_xlsx_archive(menu_book(["10", "A", "1", "Dolma", 8])) == []


@pytest.mark.parametrize(
    "data",
    [
        b"category_external_id,name",
        b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1",  # .xls / зашифрованная книга (OLE)
        b"PK\x03\x04" + b"\x00" * 64,  # повреждённый zip
        b"",
    ],
)
def test_not_a_zip_is_unsupported_format(data: bytes) -> None:
    issues = check_xlsx_archive(data)
    assert [(issue.code, issue.message) for issue in issues] == [(IssueCode.UNSUPPORTED_FORMAT, NOT_XLSX_HINT)]


def test_zip_without_workbook_is_unsupported_format() -> None:
    data = zip_of({"word/document.xml": b"<w/>", "[Content_Types].xml": b"<Types/>"})
    assert [issue.message for issue in check_xlsx_archive(data)] == [NOT_XLSX_HINT]


@pytest.mark.parametrize("name", ["xl/vbaProject.bin", "xl/VBAPROJECT.BIN", "custom/vbaProject.bin"])
def test_macros_are_rejected(name: str) -> None:
    data = rewrite(menu_book(["10", "A", "1", "Dolma", 8]), lambda _, content: content, {name: b"\x00" * 16})
    result = parse(data)
    issue = only_issue(result)
    assert (issue.code, issue.message) == (IssueCode.UNSUPPORTED_FORMAT, MACRO_HINT)
    assert result.file_format is FileFormat.XLSX


def test_zip_bomb_by_compression_ratio() -> None:
    data = zip_of({"xl/workbook.xml": b"<workbook/>", "xl/worksheets/sheet1.xml": b"\x00" * (2 * 1024 * 1024)})
    assert len(data) < 100 * 1024
    issue = only_issue(parse(data))
    assert issue.code is IssueCode.UNSAFE_ARCHIVE
    assert str(MAX_COMPRESSION_RATIO) in issue.message


def test_compression_ratio_limit_is_inclusive(monkeypatch: pytest.MonkeyPatch) -> None:
    data = menu_book(["10", "A", "1", "Dolma", 8])
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        worst = max(math.ceil(m.file_size / max(m.compress_size, 1)) for m in archive.infolist())
    monkeypatch.setattr(xlsx_reader, "MAX_COMPRESSION_RATIO", worst)
    assert check_xlsx_archive(data) == []
    monkeypatch.setattr(xlsx_reader, "MAX_COMPRESSION_RATIO", worst - 1)
    assert [issue.code for issue in check_xlsx_archive(data)] == [IssueCode.UNSAFE_ARCHIVE]


def test_uncompressed_total_over_limit(monkeypatch: pytest.MonkeyPatch) -> None:
    data = menu_book(["10", "A", "1", "Dolma", 8])
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        total = sum(member.file_size for member in archive.infolist())
    monkeypatch.setattr(xlsx_reader, "MAX_UNCOMPRESSED_BYTES", total - 1)
    issue = only_issue(parse(data))
    assert issue.code is IssueCode.UNSAFE_ARCHIVE
    monkeypatch.setattr(xlsx_reader, "MAX_UNCOMPRESSED_BYTES", total)
    assert check_xlsx_archive(data) == []


def with_part(name: str, content: bytes) -> bytes:
    return rewrite(menu_book(["10", "A", "1", "Dolma", 8]), lambda _, original: original, {name: content})


def test_non_sheet_parts_over_byte_limit(monkeypatch: pytest.MonkeyPatch) -> None:
    """Security review #65 [LOW]: общие строки и стили читают обе книги; имя части может быть любым."""
    data = with_part("xl/strings-renamed.xml", b"<sst/>" * 20)
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        size = sum(m.file_size for m in archive.infolist() if not m.filename.startswith("xl/worksheets/"))
    monkeypatch.setattr(xlsx_reader, "MAX_NON_SHEET_BYTES", size)
    assert check_xlsx_archive(data) == []
    monkeypatch.setattr(xlsx_reader, "MAX_NON_SHEET_BYTES", size - 1)
    issue = only_issue(parse(data))
    assert issue.code is IssueCode.UNSAFE_ARCHIVE
    assert "кроме листов" in issue.message


def test_non_sheet_parts_over_tag_limit() -> None:
    items = b"".join(f"<si><t>{n}</t></si>".encode() for n in range(xlsx_reader.MAX_NON_SHEET_TAGS // 2))
    strings = b"<sst>" + items + b"</sst>"
    result, elapsed = timed_parse(with_part("xl/sharedStrings.xml", strings))
    issue = only_issue(result)
    assert issue.code is IssueCode.UNSAFE_ARCHIVE
    limit = xlsx_reader.MAX_NON_SHEET_TAGS
    assert issue.message.startswith(f"служебные части книги (кроме листов) содержит больше {limit}")
    assert elapsed < 5


def test_book_over_tag_limit_counts_other_sheets(monkeypatch: pytest.MonkeyPatch) -> None:
    data = with_part("xl/worksheets/sheet9.xml", b"<worksheet>" + b"<col/>" * 50 + b"</worksheet>")
    monkeypatch.setattr(xlsx_reader, "MAX_BOOK_TAGS", 40)
    issue = only_issue(parse(data))
    assert issue.code is IssueCode.UNSAFE_ARCHIVE
    assert issue.message.startswith("книга содержит больше 40 элементов")


def test_book_budget_smaller_than_non_sheet_budget_reports_book(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(xlsx_reader, "MAX_BOOK_TAGS", 5)
    issue = only_issue(parse(menu_book(["10", "A", "1", "Dolma", 8])))
    assert issue.message.startswith("книга содержит больше 5 элементов")


def test_sheet_and_tag_limits() -> None:
    assert xlsx_reader.MAX_NON_SHEET_BYTES == 20 * 1024 * 1024
    assert xlsx_reader.MAX_SHEET_CELLS == 30 * xlsx_reader.MAX_DATA_ROWS
    assert xlsx_reader.MAX_BOOK_TAGS == xlsx_reader.MAX_SHEET_TAGS + xlsx_reader.MAX_NON_SHEET_TAGS


def test_real_size_limit_is_50_mb() -> None:
    assert xlsx_reader.MAX_UNCOMPRESSED_BYTES == 50 * 1024 * 1024
    assert MAX_COMPRESSION_RATIO == 100


# ---------------------------------------------------------------- разбор: повреждённые книги и XML


def test_damaged_workbook_xml_is_unreadable() -> None:
    data = rewrite(
        menu_book(["10", "A", "1", "Dolma", 8]),
        lambda name, content: b"<not-xml" if name == "xl/workbook.xml" else content,
    )
    assert check_xlsx_archive(data) == []
    issue = only_issue(parse(data))
    assert (issue.code, issue.message) == (IssueCode.UNSUPPORTED_FORMAT, UNREADABLE_HINT)


@pytest.mark.parametrize(
    "doctype",
    [
        b'<!DOCTYPE x [<!ENTITY a "aaaaaaaaaa"><!ENTITY b "&a;&a;&a;&a;&a;&a;&a;&a;&a;&a;">]>',  # XML-бомба
        b'<!DOCTYPE x [<!ENTITY ext SYSTEM "file:///etc/hostname">]>',  # XXE
    ],
)
def test_xml_entities_are_rejected(doctype: bytes) -> None:
    def inject(name: str, content: bytes) -> bytes:
        if name != "xl/worksheets/sheet1.xml":
            return content
        return doctype + content.split(b"?>", 1)[-1] if content.startswith(b"<?xml") else doctype + content

    data = rewrite(menu_book(["10", "A", "1", "Dolma", 8]), inject)
    # Причина — именно defusedxml (сущности запрещены), а не случайная ошибка разметки.
    with pytest.raises(ValueError) as caught:
        openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    assert isinstance(caught.value.__cause__, DefusedXmlException)
    issue = only_issue(parse(data))
    assert (issue.code, issue.message) == (IssueCode.UNSUPPORTED_FORMAT, UNREADABLE_HINT)


class _Tracker:
    """Обёртка ``load_workbook``: запоминает открытые книги, чтобы проверить их закрытие."""

    def __init__(self) -> None:
        self.books: list[Any] = []
        self.closed = 0

    def __call__(self, *args: Any, **kwargs: Any) -> Any:
        workbook = openpyxl.load_workbook(*args, **kwargs)
        original_close = workbook.close

        def close() -> None:
            self.closed += 1
            original_close()

        workbook.close = close
        self.books.append(workbook)
        return workbook


def test_workbooks_are_opened_read_only_and_closed(monkeypatch: pytest.MonkeyPatch) -> None:
    tracker = _Tracker()
    monkeypatch.setattr(xlsx_reader, "load_workbook", tracker)
    assert item_ids(parse(menu_book(["10", "A", "1", "Dolma", 8]))) == ["1"]
    # Формул на листе нет — вторая книга не открывается.
    assert [(wb.read_only, wb.data_only) for wb in tracker.books] == [(True, True)]
    assert tracker.closed == 1
    data = with_cached_value(menu_book(["10", "A", "1", "Dolma", "=4+4"]), "4+4", "8")
    assert item_ids(parse(data)) == ["1"]
    assert [(wb.read_only, wb.data_only) for wb in tracker.books[1:]] == [(True, True), (True, False)]
    assert tracker.closed == 3


def test_workbook_closed_when_sheet_not_found(monkeypatch: pytest.MonkeyPatch) -> None:
    tracker = _Tracker()
    monkeypatch.setattr(xlsx_reader, "load_workbook", tracker)
    assert codes(parse(two_sheets(), sheet="nope")) == [IssueCode.SHEET_NOT_FOUND]
    assert (len(tracker.books), tracker.closed) == (1, 1)


def test_workbooks_closed_on_error_while_reading(monkeypatch: pytest.MonkeyPatch) -> None:
    tracker = _Tracker()
    monkeypatch.setattr(xlsx_reader, "load_workbook", tracker)

    def fail(*_: Any) -> Any:
        raise KeyError("broken part")

    monkeypatch.setattr(xlsx_reader, "_read_sheet", fail)
    issue = only_issue(parse(menu_book(["10", "A", "1", "Dolma", "=1+1"])))
    assert issue.message == UNREADABLE_HINT
    assert (len(tracker.books), tracker.closed) == (2, 2)


def test_error_before_any_workbook_opened(monkeypatch: pytest.MonkeyPatch) -> None:
    def fail(*_: Any, **__: Any) -> Any:
        raise OSError("boom")

    monkeypatch.setattr(xlsx_reader, "load_workbook", fail)
    assert read_xlsx_table(menu_book(["10", "A", "1", "Dolma", 8])) == (
        None,
        [SyncIssue(IssueCode.UNSUPPORTED_FORMAT, UNREADABLE_HINT)],
    )
